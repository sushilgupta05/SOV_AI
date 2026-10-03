"""Generate synthetic SOV fixtures covering the structural edge cases in fr1_claude.md section 6.

Run directly to write them to tests/fixtures/generated/ for manual UI testing:
    python tests/fixtures/make_fixtures.py
"""
from __future__ import annotations

import random
import sys
from pathlib import Path

from openpyxl import Workbook

HEADERS = ["Loc #", "Street Address", "City", "State", "Zip", "County", "Bldg Repl Cost", "Contents",
           "BI", "Occupancy", "Construction", "# Stories", "Year Built", "Sprinklered"]
CITIES = [("Austin", "TX", "Travis"), ("Dallas", "TX", "Dallas"), ("Miami", "FL", "Miami-Dade"),
          ("Tampa", "FL", "Hillsborough"), ("Denver", "CO", "Denver")]
OCC = ["Apartment", "Office", "Retail", "Warehouse", "Light Manufacturing"]
CONST = ["Frame", "Joisted Masonry", "Non-Combustible", "Masonry Non-Combustible", "Fire Resistive"]


def sov_rows(n: int, seed: int = 7) -> list[list[object]]:
    rnd = random.Random(seed)
    rows = []
    for i in range(1, n + 1):
        city, state, county = rnd.choice(CITIES)
        rows.append([
            i, f"{rnd.randint(100, 9999)} {rnd.choice(['Oak', 'Elm', 'Main', 'Pine'])} St", city, state,
            rnd.randint(10000, 99999), county, float(rnd.randint(500, 20000) * 1000),
            float(rnd.randint(50, 2000) * 1000), float(rnd.randint(0, 500) * 1000), rnd.choice(OCC),
            rnd.choice(CONST), rnd.randint(1, 12), rnd.randint(1950, 2022), rnd.choice(["Y", "N", "Y13"]),
        ])
    return rows


def _write(ws, start_row: int, rows: list[list[object]]) -> None:
    for r, row in enumerate(rows, start=start_row):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)


def clean(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    _write(ws, 1, [HEADERS] + sov_rows(30))
    wb.save(path); return path


def title_block(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Schedule"
    ws["A1"] = "Statement of Values - Redwood Housing LLC"; ws.merge_cells("A1:N1")
    ws["A2"] = "Policy: SAMPLE-UMR-0001"
    ws["A4"] = "Prepared by [REDACTED]"; ws["B4"] = "Broker"
    ws["A5"] = "As of 2024-01-01"
    _write(ws, 6, [HEADERS] + sov_rows(25))
    wb.save(path); return path


def two_row_header(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Locations"
    ws["A1"] = "Property Schedule"; ws.merge_cells("A1:N1")
    ws["A2"] = "Location"; ws.merge_cells("A2:F2")
    ws["G2"] = "Values"; ws.merge_cells("G2:I2")
    ws["J2"] = "Construction Details"; ws.merge_cells("J2:N2")
    headers = HEADERS[:6] + ["Building", "Contents", "BI"] + HEADERS[9:]
    _write(ws, 3, [headers] + sov_rows(20))
    wb.save(path); return path


def text_heavy(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "Data"
    ws["A1"] = "Client submission"
    headers = ["Loc", "Address", "City", "State", "County", "Occupancy", "Construction", "Sprinklered",
               "Bldg Value"]
    rnd = random.Random(3)
    rows = []
    for i in range(15):
        city, state, county = rnd.choice(CITIES)
        rows.append([f"L-{i:03d}", f"{rnd.randint(1, 999)} Main Street", city, state, county,
                     rnd.choice(OCC), rnd.choice(CONST), rnd.choice(["Yes", "No"]),
                     float(rnd.randint(1, 50) * 100000)])
    _write(ws, 3, [headers] + rows)
    wb.save(path); return path


def multi_sheet(path: Path) -> Path:
    wb = Workbook()
    cover = wb.active; cover.title = "Cover"
    cover["A1"] = "Statement of Values submission"
    cover["A3"] = "Please contact the broker with any questions about this schedule."
    cover["A4"] = "All values in USD. Values as of renewal date."
    loc = wb.create_sheet("Locations")
    loc["A1"] = "Location Schedule"; loc["A2"] = "Insured: Redwood Housing LLC"
    _write(loc, 4, [HEADERS] + sov_rows(40))
    lookup = wb.create_sheet("Lookup")
    _write(lookup, 1, [["Code", "Description"]] + [[c, f"ISO class {i}"] for i, c in enumerate(CONST)])
    lookup.sheet_state = "hidden"
    wb.create_sheet("Empty")
    wb.save(path); return path


def totals(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    rows = sov_rows(12)
    _write(ws, 1, [HEADERS] + rows)
    n = len(rows) + 2
    ws.cell(row=n, column=1, value="Subtotal TX")
    ws.cell(row=n + 1, column=1, value="Grand Total")
    ws.cell(row=n + 1, column=7, value=sum(r[6] for r in rows))
    wb.save(path); return path


def data_merged(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    _write(ws, 1, [HEADERS] + sov_rows(10))
    ws["C5"] = "Austin"; ws["C6"] = None; ws["C7"] = None
    ws.merge_cells("C5:C7")
    wb.save(path); return path


def dup_blank_headers(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    headers = ["Loc #", "Address", "City", "Value", "Value", None, None, "Year Built"]
    rows = [[i, f"{i} Main St", "Austin", 1000.0 * i, 500.0 * i, f"note {i}", None, 1990 + i]
            for i in range(1, 11)]
    _write(ws, 1, [headers] + rows)
    wb.save(path); return path


def zip_text(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    rows = sov_rows(5)
    rows[0][4] = "02134"
    _write(ws, 1, [HEADERS] + rows)
    wb.save(path); return path


def preamble_csv(path: Path) -> Path:
    lines = ["Statement of Values;;", "Insured: Société Générale Immobilier;;",
             ";".join(HEADERS)]
    for r in sov_rows(20):
        lines.append(";".join("" if v is None else str(v) for v in r))
    path.write_bytes("\r\n".join(lines).encode("cp1252"))
    return path


def phantom(path: Path) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "SOV"
    _write(ws, 1, [HEADERS] + sov_rows(50))
    ws.cell(row=1_048_576, column=1).number_format = "0.00"  # formatting far beyond the data
    wb.save(path); return path


def large(path: Path, rows: int = 5000, cols: int = 40) -> Path:
    wb = Workbook(write_only=True); ws = wb.create_sheet("SOV")
    extra = [f"Extra Field {i}" for i in range(cols - len(HEADERS))]
    ws.append(HEADERS + extra)
    for r in sov_rows(rows):
        ws.append(r + [float(i) for i in range(len(extra))])
    wb.save(path); return path


def bad_files(directory: Path) -> dict[str, Path]:
    files = {
        "empty": directory / "empty.xlsx",
        "txt": directory / "notes.txt",
        "corrupt": directory / "corrupt.xlsx",
        "legacy": directory / "legacy.xlsx",
    }
    files["empty"].write_bytes(b"")
    files["txt"].write_text("hello")
    files["corrupt"].write_bytes(b"PK\x03\x04" + b"\x00garbage" * 50)
    files["legacy"].write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    return files


BUILDERS = {
    "clean.xlsx": clean, "title_block.xlsx": title_block, "two_row_header.xlsx": two_row_header,
    "text_heavy.xlsx": text_heavy, "multi_sheet.xlsx": multi_sheet, "totals.xlsx": totals,
    "data_merged.xlsx": data_merged, "dup_blank_headers.xlsx": dup_blank_headers,
    "zip_text.xlsx": zip_text, "preamble.csv": preamble_csv, "phantom.xlsx": phantom,
}


def build_all(directory: Path, include_large: bool = True) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    out = {name: fn(directory / name) for name, fn in BUILDERS.items()}
    if include_large:
        out["large.xlsx"] = large(directory / "large.xlsx")
    out.update({f"bad_{k}": v for k, v in bad_files(directory).items()})
    return out


if __name__ == "__main__":
    target = Path(__file__).parent / "generated"
    for name, p in build_all(target, include_large="--no-large" not in sys.argv).items():
        print(f"{name:28s} {p}")
